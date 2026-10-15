"""xlsx_loader.py – every worksheet row becomes one sentence.

Example: "Sheet Sales | Region=South | Q3=1.2M"  (row number = real Excel row).
"""
from pathlib import Path

from openpyxl import load_workbook


def _fmt(v) -> str:
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def load_xlsx(path: str, **_) -> list[dict]:
    wb = load_workbook(path, data_only=True, read_only=True)
    name = Path(path).name
    out: list[dict] = []
    for ws in wb.worksheets:
        header: list[str] | None = None
        for row_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if not any(c is not None and str(c).strip() for c in row):
                continue
            if header is None:  # first non-empty row = header
                header = [_fmt(c) if c is not None else f"col{i+1}" for i, c in enumerate(row)]
                continue
            parts = [
                f"{header[i] if i < len(header) else f'col{i+1}'}={_fmt(c)}"
                for i, c in enumerate(row)
                if c is not None and str(c).strip()
            ]
            out.append(
                {
                    "text": f"Sheet {ws.title} | " + " | ".join(parts),
                    "page_or_row": f"row {row_no}",
                    "metadata": {"kind": "row", "sheet": ws.title, "row": row_no, "file": name},
                }
            )
    wb.close()
    return out
