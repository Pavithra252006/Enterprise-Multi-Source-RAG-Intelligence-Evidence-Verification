"""
generate_ppt.py - Generates an 8-slide widescreen PPTX matching the SYNAPSE '26 Hackathon format.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ── Color Palette ─────────────────────────────────────────────────────────────
BG_COLOR = RGBColor(248, 250, 252)         # #F8FAFC
DARK_NAVY = RGBColor(15, 23, 42)           # #0F172A
TEAL_CYAN = RGBColor(13, 148, 136)         # #0D9488
CARD_BG = RGBColor(241, 245, 249)          # #F1F5F9
CARD_BORDER = RGBColor(226, 232, 240)      # #E2E8F0
TEXT_MUTED = RGBColor(100, 116, 139)       # #64748B
TEXT_DARK = RGBColor(15, 23, 42)           # #0F172A
WHITE = RGBColor(255, 255, 255)
ACCENT_BLUE = RGBColor(14, 116, 144)       # #0E7490
PURPLE_ACCENT = RGBColor(126, 34, 206)     # #7E22CE

def apply_header(slide, title_text):
    # Top Tag
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(0.4))
    tf = tag_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "S Y N A P S E  ' 2 6   |   2 4 - H O U R   H A C K A T H O N"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = TEAL_CYAN

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11), Inches(0.6))
    tf2 = title_box.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = title_text
    p2.font.name = "Arial Black"
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = DARK_NAVY

def apply_footer(slide, page_num_str):
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11), Inches(0.35))
    tf = footer_box.text_frame
    p = tf.paragraphs[0]
    p.text = f"V.S.B. Engineering College • {page_num_str}"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.color.rgb = TEXT_MUTED

def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# ==============================================================================
# SLIDE 1: TEAM & PROJECT DETAILS
# ==============================================================================
slide1 = prs.slides.add_slide(blank_layout)
set_slide_background(slide1)
apply_header(slide1, "TEAM & PROJECT DETAILS")
apply_footer(slide1, "02")

rows = [
    ("Project Title", "PS-019 Enterprise Multi-Source RAG Intelligence & Evidence Verification"),
    ("Team Name / Team ID", "ERROR 404 (Team ID: SY05)\nMembers: Pavithra S, Nithya Sri S"),
    ("Problem Statement ID and Title", "PS-019 — Enterprise Multi-Source RAG Intelligence & Evidence Verification"),
    ("Team Leader", "Pavithra S — [922524243142]; AI & DS, III Year, Section [C]"),
    ("Project Repository / Demo Link", "Repo: github.com/Pavithra252006/Enterprise-Multi-Source-RAG-Intelligence-Evidence-Verification\nDemo Link: http://localhost:5173  (Network: http://10.150.58.249:5173)"),
]

top_pos = 1.6
for label, val in rows:
    # Left Label Box
    sh1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top_pos), Inches(3.6), Inches(0.85))
    sh1.fill.solid()
    sh1.fill.fore_color.rgb = CARD_BG
    sh1.line.color.rgb = CARD_BORDER
    sh1.line.width = Pt(1)
    tf1 = sh1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = label
    p1.font.bold = True
    p1.font.size = Pt(13)
    p1.font.color.rgb = ACCENT_BLUE

    # Right Value Box
    sh2 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.6), Inches(top_pos), Inches(7.9), Inches(0.85))
    sh2.fill.solid()
    sh2.fill.fore_color.rgb = WHITE
    sh2.line.color.rgb = CARD_BORDER
    sh2.line.width = Pt(1)
    tf2 = sh2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = val
    p2.font.size = Pt(12)
    p2.font.color.rgb = DARK_NAVY
    if "\n" in val:
        for idx, line in enumerate(val.split("\n")[1:], start=1):
            p_sub = tf2.add_paragraph()
            p_sub.text = line
            p_sub.font.size = Pt(11)
            p_sub.font.color.rgb = TEXT_MUTED

    top_pos += 1.02

# ==============================================================================
# SLIDE 2: PROBLEM STATEMENT
# ==============================================================================
slide2 = prs.slides.add_slide(blank_layout)
set_slide_background(slide2)
apply_header(slide2, "PROBLEM STATEMENT")
apply_footer(slide2, "03")

prob_rows = [
    ("PROBLEM STATEMENT", "Build an advanced RAG platform that answers complex questions over multiple enterprise sources (PDF, HTML, CSV, XLSX) with evidence verification, conflict resolution, abstention, access control, and What-If counterfactual policy simulation."),
    ("PROBLEM FACED", "Employees get fragmented, outdated or leaked answers: tables lose structure, restricted data reaches low-clearance users, old policy versions contradict new ones, LLMs hallucinate without citations, and leadership cannot forecast the downstream impact of policy changes."),
    ("SUCCESS OUTCOME", "Zero access-control leakage, 100% citation correctness, <5% unsupported claims, 100% temporal conflict resolution, and an automated What-If Policy Impact Simulator forecasting downstream contradictions and amendments."),
]

top_pos = 1.7
for label, desc in prob_rows:
    # Left dark box
    sh1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top_pos), Inches(2.8), Inches(1.4))
    sh1.fill.solid()
    sh1.fill.fore_color.rgb = DARK_NAVY
    sh1.line.color.rgb = DARK_NAVY
    tf1 = sh1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = label
    p1.font.bold = True
    p1.font.size = Pt(12)
    p1.font.color.rgb = WHITE

    # Right light box
    sh2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.8), Inches(top_pos), Inches(8.7), Inches(1.4))
    sh2.fill.solid()
    sh2.fill.fore_color.rgb = CARD_BG
    sh2.line.color.rgb = CARD_BORDER
    sh2.line.width = Pt(1)
    tf2 = sh2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = desc
    p2.font.size = Pt(13)
    p2.font.color.rgb = DARK_NAVY

    top_pos += 1.65

# ==============================================================================
# SLIDE 3: PROPOSED SOLUTION
# ==============================================================================
slide3 = prs.slides.add_slide(blank_layout)
set_slide_background(slide3)
apply_header(slide3, "PROPOSED SOLUTION")
apply_footer(slide3, "04")

# Top summary box
sum_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.85))
sum_box.fill.solid()
sum_box.fill.fore_color.rgb = CARD_BG
sum_box.line.color.rgb = CARD_BORDER
tf = sum_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Solution summary: Clearance-gated hybrid retrieval, neural reranking, claim-level verification, and a What-If Policy Impact Simulator produce verified answers, abstain on weak evidence, and forecast corporate policy risks."
p.font.size = Pt(12)
p.font.color.rgb = DARK_NAVY

# Target users
p_sub = tf.add_paragraph()
p_sub.text = "Target users: Interns, analysts, managers, executives (4-tier clearance), and corporate governance/compliance teams."
p_sub.font.bold = True
p_sub.font.size = Pt(12)
p_sub.font.color.rgb = TEAL_CYAN

# 4 Core Features Cards
features = [
    ("Core feature 1", "Access-gated hybrid retrieval: FAISS + BM25 with RRF (k=60), then bge cross-encoder rerank."),
    ("Core feature 2", "Claim-level strict verification with inline citations and fail-safe abstention (<0.60)."),
    ("Core feature 3", "Temporal conflict resolution: newer policy versions (v2 2024) supersede older ones (v1 2022)."),
    ("Core feature 4 (NEW)", "What-If Policy Simulator: counterfactual engine predicting downstream contradictions and amendments."),
]

left_pos = 0.8
for title, desc in features:
    sh = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left_pos), Inches(2.6), Inches(2.75), Inches(2.6))
    sh.fill.solid()
    sh.fill.fore_color.rgb = WHITE
    sh.line.color.rgb = CARD_BORDER
    sh.line.width = Pt(1)
    tf = sh.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.bold = True
    p1.font.size = Pt(13)
    p1.font.color.rgb = PURPLE_ACCENT if "NEW" in title else ACCENT_BLUE

    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(11)
    p2.font.color.rgb = DARK_NAVY
    left_pos += 2.98

# Bottom Dark Banner
bot_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.45), Inches(11.7), Inches(1.2))
bot_box.fill.solid()
bot_box.fill.fore_color.rgb = DARK_NAVY
bot_box.line.color.rgb = DARK_NAVY
tf_b = bot_box.text_frame
tf_b.word_wrap = True
pb = tf_b.paragraphs[0]
pb.text = "What is different? Access control executes strictly before reranking and LLM synthesis (0-leakage guarantee), and the platform includes counterfactual policy simulation to stress-test corporate decisions before rollout."
pb.font.size = Pt(13)
pb.font.color.rgb = WHITE

# ==============================================================================
# SLIDE 4: SOLUTION WORKFLOW
# ==============================================================================
slide4 = prs.slides.add_slide(blank_layout)
set_slide_background(slide4)
apply_header(slide4, "SOLUTION WORKFLOW")
apply_footer(slide4, "05")

flow_steps = [
    ("USER INPUT", "Query / Policy Scenario + X-User-Id (role/clearance)"),
    ("PROCESSING / AI LOGIC", "Clearance filter, Dense + BM25 top-30, RRF, bge rerank"),
    ("BACKEND / SERVICES", "FastAPI: Claude generation, claim verifier, What-If simulator"),
    ("DATABASE / SOURCES", "FAISS, BM25, SQL; PDF, HTML, CSV, XLSX corpus"),
    ("OUTPUT / ACTION", "Verified answer with citations, abstention, or Policy Risk Matrix"),
]

left_pos = 0.8
for i, (title, desc) in enumerate(flow_steps):
    sh = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left_pos), Inches(1.6), Inches(2.1), Inches(2.6))
    sh.fill.solid()
    sh.fill.fore_color.rgb = DARK_NAVY if i in [0, 4] else CARD_BG
    sh.line.color.rgb = CARD_BORDER
    sh.line.width = Pt(1)
    tf = sh.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.bold = True
    p1.font.size = Pt(11)
    p1.font.color.rgb = WHITE if i in [0, 4] else ACCENT_BLUE

    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(11)
    p2.font.color.rgb = WHITE if i in [0, 4] else DARK_NAVY

    left_pos += 2.4

# Bottom Key Decisions Box
key_box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.5), Inches(11.7), Inches(2.1))
key_box.fill.solid()
key_box.fill.fore_color.rgb = WHITE
key_box.line.color.rgb = CARD_BORDER
key_box.line.width = Pt(1)
tf_k = key_box.text_frame
tf_k.word_wrap = True
pk = tf_k.paragraphs[0]
pk.text = "Key decisions / validation / security"
pk.font.bold = True
pk.font.size = Pt(13)
pk.font.color.rgb = TEAL_CYAN

bullets = [
    "• Clearance filter runs before retrieval scoring: 0-leakage by architectural design.",
    "• Rerank score below 0.40 triggers adaptive multi-round query decomposition (max 2 rounds).",
    "• Claims must be strictly SUPPORTED by verified chunks; confidence below 0.60 forces fail-safe abstention.",
    "• What-If Policy Engine executes multi-source dependency traversal, forecasting legal, financial, and operational clashes."
]
for b in bullets:
    p_b = tf_k.add_paragraph()
    p_b.text = b
    p_b.font.size = Pt(12)
    p_b.font.color.rgb = DARK_NAVY

# ==============================================================================
# SLIDE 5: SYSTEM ARCHITECTURE
# ==============================================================================
slide5 = prs.slides.add_slide(blank_layout)
set_slide_background(slide5)
apply_header(slide5, "SYSTEM ARCHITECTURE")
apply_footer(slide5, "06")

arch_blocks = [
    ("FRONTEND", "React 19 + Vite + Tailwind v4: Q&A Chat, Policy Simulator Lab, Benchmark Eval"),
    ("BACKEND / API", "FastAPI (async Python 3.12): Retrieval, Verifier, and Policy Simulation routers"),
    ("DATABASE", "FAISS IndexFlatIP, BM25 index, SQLAlchemy 2.0 Async (SQLite / PostgreSQL)"),
    ("EXTERNAL SERVICES", "Claude 3.5 Sonnet, bge-reranker-base, bge-small-en-v1.5 embeddings"),
]

left_pos = 0.8
for title, desc in arch_blocks:
    sh = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left_pos), Inches(1.6), Inches(2.7), Inches(1.9))
    sh.fill.solid()
    sh.fill.fore_color.rgb = DARK_NAVY if "FRONTEND" in title else (ACCENT_BLUE if "API" in title else CARD_BG)
    sh.line.color.rgb = CARD_BORDER
    sh.line.width = Pt(1)
    tf = sh.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.bold = True
    p1.font.size = Pt(13)
    p1.font.color.rgb = WHITE if ("FRONTEND" in title or "API" in title) else DARK_NAVY

    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(11)
    p2.font.color.rgb = WHITE if ("FRONTEND" in title or "API" in title) else DARK_NAVY
    left_pos += 3.0

# Architecture Bullets
arch_box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.8), Inches(11.7), Inches(2.8))
arch_box.fill.solid()
arch_box.fill.fore_color.rgb = WHITE
arch_box.line.color.rgb = CARD_BORDER
arch_box.line.width = Pt(1)
tfa = arch_box.text_frame
tfa.word_wrap = True
arch_bullets = [
    "• Responsibility: UI provides Q&A Chat and What-If Policy Lab; API executes access gating, hybrid retrieval, verification & simulation.",
    "• Data flow: REST call ➔ clearance filter ➔ hybrid RRF search ➔ rerank ➔ Claude/Deterministic engine ➔ claim verifier ➔ UI Card.",
    "• Simulation Pipeline: Takes proposed policy change ➔ semantic dependency search ➔ contradiction detection ➔ amendment generation.",
    "• Ingestion: Structure-aware chunking with metadata feeds FAISS, BM25 and SQL across PDF, HTML, CSV, and XLSX.",
    "• Failure handling: Built-in deterministic fallback engine for zero-API-cost offline demo; fail-safe abstention on low evidence."
]
for i, b in enumerate(arch_bullets):
    p = tfa.paragraphs[0] if i == 0 else tfa.add_paragraph()
    p.text = b
    p.font.size = Pt(12)
    p.font.color.rgb = DARK_NAVY

# ==============================================================================
# SLIDE 6: TECHNOLOGY STACK
# ==============================================================================
slide6 = prs.slides.add_slide(blank_layout)
set_slide_background(slide6)
apply_header(slide6, "TECHNOLOGY STACK")
apply_footer(slide6, "09")

tech_cards = [
    ("Frontend", "React 19, Vite 8, Tailwind CSS v4: fast, component-based answer cards and What-If Policy Lab."),
    ("Backend / APIs", "FastAPI + Uvicorn: async, typed, low-latency endpoints with strict Pydantic validation."),
    ("Database / Storage", "FAISS IndexFlatIP, BM25 index, SQLAlchemy 2.0 Async + Alembic (SQLite / PostgreSQL)."),
    ("AI / ML / Algorithms", "bge-small-en-v1.5 embeddings, Reciprocal Rank Fusion (k=60), bge-reranker-base, Claude 3.5 Sonnet."),
    ("Deployment / Hosting", "Production Docker container, Render Blueprint (render.yaml), Vercel edge SPA routing."),
    ("External APIs / Datasets", "Anthropic Claude API; synthetic enterprise corpus across PDF, HTML wiki, CSV, and Excel spreadsheets."),
]

positions = [
    (0.8, 1.6), (6.8, 1.6),
    (0.8, 3.35), (6.8, 3.35),
    (0.8, 5.1), (6.8, 5.1),
]

for (title, desc), (l, t) in zip(tech_cards, positions):
    sh = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(5.7), Inches(1.5))
    sh.fill.solid()
    sh.fill.fore_color.rgb = CARD_BG
    sh.line.color.rgb = CARD_BORDER
    sh.line.width = Pt(1)
    tf = sh.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.bold = True
    p1.font.size = Pt(13)
    p1.font.color.rgb = ACCENT_BLUE

    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(11)
    p2.font.color.rgb = DARK_NAVY

# ==============================================================================
# SLIDE 7: CORE FEATURE DEMONSTRATION
# ==============================================================================
slide7 = prs.slides.add_slide(blank_layout)
set_slide_background(slide7)
apply_header(slide7, "CORE FEATURE DEMONSTRATION")
apply_footer(slide7, "11")

demo_flow = [
    ("INPUT", "1. Query: Remote work policy, as Intern vs Executive\n2. Simulation: 4-Day Remote Work Policy Proposal"),
    ("PROCESS", "Clearance filter ➔ hybrid retrieval ➔ rerank ➔ conflict check ➔ claim verifier & What-If engine"),
    ("OUTPUT", "1. Cited v2.0 rule (3 days); v1.0 superseded; salary hidden from intern\n2. Policy Risk Matrix + Clashing Clauses + Amendments"),
]

left_pos = 0.8
for i, (title, desc) in enumerate(demo_flow):
    sh = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left_pos), Inches(1.6), Inches(3.7), Inches(2.0))
    sh.fill.solid()
    sh.fill.fore_color.rgb = DARK_NAVY if i == 0 else (ACCENT_BLUE if i == 1 else DARK_NAVY)
    sh.line.color.rgb = CARD_BORDER
    sh.line.width = Pt(1)
    tf = sh.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.bold = True
    p1.font.size = Pt(13)
    p1.font.color.rgb = WHITE

    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(11)
    p2.font.color.rgb = WHITE
    left_pos += 4.0

# Bottom Details Box
detail_box = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.9), Inches(11.7), Inches(2.7))
detail_box.fill.solid()
detail_box.fill.fore_color.rgb = WHITE
detail_box.line.color.rgb = CARD_BORDER
detail_box.line.width = Pt(1)
tfd = detail_box.text_frame
tfd.word_wrap = True

demos = [
    ("Test scenario 1 (Conflict & Clearance): ", "Conflicting 2022 (v1.0, 2 days) and 2024 (v2.0, 3 days) remote-work policies."),
    ("Expected output: ", "Current rule (3 days) with inline citation; restricted compensation & vendor pricing hidden from low roles."),
    ("Test scenario 2 (What-If Simulation): ", "Proposal to expand remote work to 4 days/week."),
    ("Simulation output: ", "Flags HIGH risk conflict with IT Security Clause 4.3 (monthly hardware token calibration), calculates -40% desk utilization shift, and outputs recommended policy amendment clauses."),
    ("Evidence: ", "20-case benchmark: 100% citation correctness, 0% unsupported claims, 0 leakage, 100% conflict resolution."),
    ("Demo link: ", "http://localhost:5173  (Network: http://10.150.58.249:5173)")
]

for i, (bold_label, text_val) in enumerate(demos):
    p = tfd.paragraphs[0] if i == 0 else tfd.add_paragraph()
    p.text = bold_label + text_val
    p.font.size = Pt(11)
    p.font.color.rgb = DARK_NAVY

# ==============================================================================
# SLIDE 8: INNOVATION, FEASIBILITY & IMPACT
# ==============================================================================
slide8 = prs.slides.add_slide(blank_layout)
set_slide_background(slide8)
apply_header(slide8, "INNOVATION, FEASIBILITY & IMPACT")
apply_footer(slide8, "13")

innov_cards = [
    ("Innovation", "Pre-retrieval clearance gating, strict claim-level judge verifier, temporal conflict supersession, and Counterfactual Policy Simulator."),
    ("Expected users / beneficiaries", "Enterprise employees, compliance and audit teams, HR leadership, and C-suite security officers."),
    ("Expected impact", "Auditable answers, strict zero-leakage access control, elimination of policy contradictions, and pre-tested policy rollout safety."),
    ("Feasibility", "Open-source stack; runs locally on CPU with zero external API dependencies (built-in fallback engine); Claude API optional."),
    ("Scalability", "Swap SQLite for PostgreSQL / Milvus; GPU reranking; horizontally scalable FastAPI workers and streaming SSE citations."),
    ("Future improvements", "Dynamic cell-level tabular masking, Slack/Teams bot integration, and Neo4j knowledge-graph relationship traversal."),
]

for (title, desc), (l, t) in zip(innov_cards, positions):
    sh = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(5.7), Inches(1.5))
    sh.fill.solid()
    sh.fill.fore_color.rgb = CARD_BG
    sh.line.color.rgb = CARD_BORDER
    sh.line.width = Pt(1)
    tf = sh.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.bold = True
    p1.font.size = Pt(13)
    p1.font.color.rgb = TEAL_CYAN if "Innovation" in title else ACCENT_BLUE

    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(11)
    p2.font.color.rgb = DARK_NAVY

# Save presentation
output_path = "PS019_Enterprise_RAG_Presentation.pptx"
prs.save(output_path)
print(f"Presentation saved successfully to {output_path}")
